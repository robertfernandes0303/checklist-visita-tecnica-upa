# -*- coding: utf-8 -*-
"""

Correcoes em relacao ao modulo institucional (licoes do incidente de 30/09):
  - Rascunho por CODIGO de visita (nao por unidade): dois auditores nao se
    sobrescrevem.
  - Autosave com numero de VERSAO e trava por visita: uma gravacao mais antiga
    nunca sobrescreve uma mais nova, mesmo terminando depois.
  - Falha de autosave NAO cai em arquivo local (que nunca seria lido de volta):
    fica registrada e a tela avisa o auditor (ver estado_autosave).
  - Gravacao com value_input_option=RAW: texto como "=1+1" nao vira formula.
  - Envio idempotente: cada bloco (Itens, Evidencias, Visitas) so e gravado se
    o codigo da visita ainda nao existe na aba; Visitas e gravada por ultimo
    (marca de conclusao). Reenviar apos uma falha nao duplica linhas.
  - Verificacao de rascunho com 3 resultados (lista / vazio / ERRO).
"""
import json
import re
from datetime import datetime
import os
import threading

from schema import (
    COLUNAS_EVIDENCIA,
    COLUNAS_ITEM,
    COLUNAS_SECAO,
    COLUNAS_VISITA,
    agora_brasilia,
    chave_hospital,
)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(BASE_DIR, "config", "mapeamento_visita_tecnica.json")
LOCAL_DIR = os.path.join(BASE_DIR, "dados_locais")
ARQUIVO_LOCAL = os.path.join(LOCAL_DIR, "Checklist_visita_tecnica_UPA.xlsx")

SCOPES = [
    "https://www.googleapis.com/auth/spreadsheets",
    "https://www.googleapis.com/auth/drive",
]
ABA_VISITAS, ABA_ITENS, ABA_EVIDENCIAS, ABA_RASCUNHO = "UPA_Visitas", "UPA_Itens", "UPA_Evidencias", "UPA_Rascunho"
ABA_SECOES = "UPA_Secoes"
# O JSON do rascunho e dividido em blocos (uma celula do Sheets aceita 50.000 caracteres)
TAM_BLOCO_RASCUNHO = 40000
BLOCOS_RASCUNHO = 12
COLUNAS_RASCUNHO = ["Código da Visita", "Unidade", "Chave da Unidade", "Responsável", "Atualizado em", "Dados (JSON)"]
PREFIXO_PLACEHOLDER = "COLOQUE_"
ERRO_VERIFICACAO_RASCUNHO = "__ERRO_VERIFICACAO_RASCUNHO__"
RAW = "RAW"

_CACHE_CLIENTE = {"cliente": None}
_CACHE_PLANILHA_ID = {}
_TRAVA_GLOBAL = threading.Lock()
_TRAVAS_VISITA = {}
_ULTIMA_VERSAO = {}
_ESTADO_AUTOSAVE = {}


# ---------------------------------------------------------------------------
# Configuracao e credenciais
# ---------------------------------------------------------------------------
def _configurado(valor):
    return bool(valor) and not str(valor).startswith(PREFIXO_PLACEHOLDER)


def _carregar_mapeamento():
    try:
        with open(CONFIG_PATH, encoding="utf-8") as f:
            return json.load(f)
    except FileNotFoundError:
        return {}


def _tem_credencial():
    if os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON"):
        return True
    try:
        import streamlit as st
        return "gcp_service_account" in st.secrets
    except Exception:
        return False


def _usa_google(mapping):
    return _tem_credencial() and (_configurado(mapping.get("planilha_id")) or _configurado(mapping.get("pasta_planilha")))


def _obter_cliente():
    if _CACHE_CLIENTE["cliente"] is not None:
        return _CACHE_CLIENTE["cliente"]
    import gspread
    from google.oauth2.service_account import Credentials

    creds_json = os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON")
    if creds_json:
        info = json.loads(creds_json)
    else:
        import streamlit as st
        info = dict(st.secrets["gcp_service_account"])
    cliente = gspread.authorize(Credentials.from_service_account_info(info, scopes=SCOPES))
    _CACHE_CLIENTE["cliente"] = cliente
    return cliente


def _config_upload_apps_script():
    """Endpoint do Apps Script que grava as fotos no Drive do dono da pasta.
    Variavel GOOGLE_UPLOAD_JSON ou secrets [upload_drive] (url, token)."""
    bruto = os.environ.get("GOOGLE_UPLOAD_JSON")
    if bruto:
        return json.loads(bruto)
    try:
        import streamlit as st
        if "upload_drive" in st.secrets:
            return dict(st.secrets["upload_drive"])
    except Exception:
        pass
    return None


def drive_para_evidencias():
    """(destino de upload, id da pasta de evidencias) ou (None, None) => modo local.
    Prefere o Apps Script (grava com o espaco do dono da pasta); sem ele usa a conta de
    servico (que so grava em Drive Compartilhado)."""
    mapping = _carregar_mapeamento()
    if not _configurado(mapping.get("pasta_evidencias")):
        return None, None
    pasta = mapping["pasta_evidencias"]
    cfg = _config_upload_apps_script()
    if cfg and cfg.get("url") and cfg.get("token"):
        from evidencias import UploadAppsScript
        return UploadAppsScript(cfg["url"], cfg["token"]), pasta
    if _tem_credencial():
        return _obter_cliente().http_client.session, pasta
    return None, None


def _resolver_planilha_id(cliente, mapping):
    if _configurado(mapping.get("planilha_id")):
        return mapping["planilha_id"]
    pasta = mapping.get("pasta_planilha")
    nome = mapping.get("nome_planilha") or "Checklist_visita_tecnica"
    chave = (pasta, nome)
    if chave in _CACHE_PLANILHA_ID:
        return _CACHE_PLANILHA_ID[chave]
    achados = cliente.list_spreadsheet_files(title=nome, folder_id=pasta)
    if achados:
        sheet_id = achados[0]["id"]
    else:
        try:
            sheet_id = cliente.create(nome, folder_id=pasta).id
        except Exception as exc:
            if "quota" in str(exc).lower():
                raise RuntimeError(
                    f"Não foi possível criar '{nome}': a conta de serviço não tem espaço próprio no Drive. "
                    "Crie a planilha manualmente, com esse nome exato, dentro da pasta (ou informe 'planilha_id')."
                ) from exc
            raise
    _CACHE_PLANILHA_ID[chave] = sheet_id
    return sheet_id


def _abrir_planilha():
    mapping = _carregar_mapeamento()
    cliente = _obter_cliente()
    return cliente.open_by_key(_resolver_planilha_id(cliente, mapping))


def _aba(planilha, titulo, colunas):
    import gspread
    try:
        return planilha.worksheet(titulo)
    except gspread.exceptions.WorksheetNotFound:
        ws = planilha.add_worksheet(title=titulo, rows=1000, cols=len(colunas) + 2)
        ws.append_row(colunas, value_input_option=RAW)
        return ws


def _aba_se_existe(planilha, titulo):
    import gspread
    try:
        return planilha.worksheet(titulo)
    except gspread.exceptions.WorksheetNotFound:
        return None


# ---------------------------------------------------------------------------
# Envio final (idempotente)
# ---------------------------------------------------------------------------
def _linhas(colunas, registros):
    return [[str(r.get(c, "")) for c in colunas] for r in registros]


def _blocos(visita, itens, evidencias, secoes=()):
    # Visitas por ULTIMO: e a marca de que o envio terminou.
    return [
        (ABA_ITENS, COLUNAS_ITEM, _linhas(COLUNAS_ITEM, itens)),
        (ABA_EVIDENCIAS, COLUNAS_EVIDENCIA, _linhas(COLUNAS_EVIDENCIA, evidencias)),
        (ABA_SECOES, COLUNAS_SECAO, _linhas(COLUNAS_SECAO, secoes)),
        (ABA_VISITAS, COLUNAS_VISITA, _linhas(COLUNAS_VISITA, [visita])),
    ]


def _gravar_google(visita, itens, evidencias, secoes=()):
    planilha = _abrir_planilha()
    codigo = str(visita["Código da Visita"])
    gravados, ja_existiam = [], []
    for titulo, colunas, linhas in _blocos(visita, itens, evidencias, secoes):
        if not linhas:
            continue
        aba = _aba(planilha, titulo, colunas)
        if codigo in aba.col_values(1):
            ja_existiam.append(titulo)
            continue
        aba.append_rows(linhas, value_input_option=RAW)
        gravados.append(titulo)
    return gravados, ja_existiam


def _estilizar_cabecalho(ws, colunas):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for i, titulo in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor="FF005F8A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 24
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _gravar_local(visita, itens, evidencias, secoes=()):
    import openpyxl
    os.makedirs(LOCAL_DIR, exist_ok=True)
    wb = openpyxl.load_workbook(ARQUIVO_LOCAL) if os.path.exists(ARQUIVO_LOCAL) else openpyxl.Workbook()
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) == 1:
        wb.remove(wb["Sheet"])
    codigo = str(visita["Código da Visita"])
    gravados, ja_existiam = [], []
    for titulo, colunas, linhas in _blocos(visita, itens, evidencias, secoes):
        if not linhas:
            continue
        if titulo not in wb.sheetnames:
            _estilizar_cabecalho(wb.create_sheet(titulo), colunas)
        ws = wb[titulo]
        if codigo in {str(c.value) for c in ws["A"][1:]}:
            ja_existiam.append(titulo)
            continue
        for linha in linhas:
            ws.append(linha)
            for cel in ws[ws.max_row]:  # texto sempre como texto (nunca formula)
                if isinstance(cel.value, str):
                    cel.data_type = "s"
        gravados.append(titulo)
    wb.save(ARQUIVO_LOCAL)
    return gravados, ja_existiam


def salvar_auditoria(visita: dict, itens: list, evidencias: list, secoes=()):
    """Retorna (sucesso, modo, mensagem). Nunca lanca excecao.
    modos: google_sheets | local_fallback | erro_google_sheets | erro_local"""
    mapping = _carregar_mapeamento()
    if _usa_google(mapping):
        try:
            gravados, existiam = _gravar_google(visita, itens, evidencias, secoes)
            msg = "Visita gravada no Google Sheets."
            if existiam and not gravados:
                msg = "Esta visita já havia sido gravada antes (nada foi duplicado)."
            elif existiam:
                msg = f"Gravação concluída; já existiam: {', '.join(existiam)}."
            return True, "google_sheets", msg
        except Exception as exc:
            try:
                _gravar_local(visita, itens, evidencias, secoes)
            except Exception:
                pass
            return False, "erro_google_sheets", f"Falha ao gravar no Google Sheets: {exc}"
    try:
        _gravar_local(visita, itens, evidencias, secoes)
        return True, "local_fallback", "Credenciais do Google não configuradas: gravado localmente em modo de teste."
    except Exception as exc:
        return False, "erro_local", f"Falha ao gravar localmente: {exc}"


# ---------------------------------------------------------------------------
# Rascunho (um por VISITA)
# ---------------------------------------------------------------------------
def _trava(codigo):
    with _TRAVA_GLOBAL:
        return _TRAVAS_VISITA.setdefault(codigo, threading.Lock())


def _caminho_rascunho_local(codigo):
    return os.path.join(LOCAL_DIR, f"Rascunho_{codigo}.json")


def _linha_do_codigo(aba, codigo):
    for i, valor in enumerate(aba.col_values(1)[1:], start=2):
        if valor == codigo:
            return i
    return None


# ---------------------------------------------------------------------------
# Mesclagem de rascunhos: duas sessoes (dois celulares, duas abas, ou uma sessao antiga que
# reabriu) podem gravar a mesma visita. Em vez de "a ultima grava por cima" (o que apagava o
# trabalho da outra), cada gravacao le o que ja esta salvo e junta os dois lados.
# ---------------------------------------------------------------------------
_MESCLADO = {}          # codigo -> dados mesclados que a sessao ainda nao conhece


def _ts(entrada):
    for chave in ("atualizado_em", "respondido_em"):
        v = (entrada or {}).get(chave)
        if v:
            try:
                return datetime.strptime(v, "%d/%m/%Y %H:%M:%S")
            except Exception:
                pass
    return None


def _chave_anexo(a):
    return (a or {}).get("id_arquivo") or (a or {}).get("link") or (a or {}).get("nome")


def _mesclar_resposta(eg, en):
    tg, tn = _ts(eg), _ts(en)
    novo_e_mais_recente = tn is not None and (tg is None or tn >= tg)
    base, outro = (en, eg) if (novo_e_mais_recente or (tg is None and tn is None)) else (eg, en)
    saida = dict(base)
    for k, v in outro.items():                     # campo vazio no lado vencedor: aproveita o outro
        if k not in saida or (saida[k] in ("", None) and v not in ("", None)):
            saida[k] = v
    return saida


def mesclar_dados(guardado, novo):
    """Junta o rascunho ja gravado com o da sessao atual, sem perder o que so um dos lados tem."""
    if not isinstance(guardado, dict):
        return novo
    saida = dict(novo)
    g, n = guardado.get("respostas") or {}, novo.get("respostas") or {}
    resp = {}
    for k in dict.fromkeys(list(g) + list(n)):
        if k not in g:
            resp[k] = n[k]
        elif k not in n:
            resp[k] = g[k]
        else:
            resp[k] = _mesclar_resposta(g[k], n[k])
    saida["respostas"] = resp
    g, n = guardado.get("previsto_informado") or {}, novo.get("previsto_informado") or {}
    prev = dict(g)
    for k, v in n.items():
        if v not in ("", None) or k not in prev:
            prev[k] = v
    saida["previsto_informado"] = prev
    removidos = list(dict.fromkeys(list(guardado.get("anexos_removidos") or []) +
                                   list(novo.get("anexos_removidos") or [])))
    ga, na = guardado.get("anexos") or {}, novo.get("anexos") or {}
    anexos = {}
    for k in dict.fromkeys(list(ga) + list(na)):
        lista, vistos = [], set()
        for a in list(ga.get(k) or []) + list(na.get(k) or []):
            ch = _chave_anexo(a)
            if ch in vistos or ch in removidos:
                continue
            vistos.add(ch)
            lista.append(a)
        if lista:
            anexos[k] = lista
    saida["anexos"] = anexos
    saida["anexos_removidos"] = removidos
    for campo in ("grupos_visitados", "grupos_concluidos"):
        saida[campo] = list(dict.fromkeys(list(guardado.get(campo) or []) + list(novo.get(campo) or [])))
    gi_, ni_ = guardado.get("impactos") or {}, novo.get("impactos") or {}
    if gi_ or ni_:
        imp = dict(gi_)
        imp.update({k: v for k, v in ni_.items() if str(v).strip() or k not in imp})
        saida["impactos"] = imp
    for k in [k for k in saida if k not in novo and k not in guardado]:
        del saida[k]                              # nao inventa campos que nenhum lado tem
    return saida


def _ler_dados_linha(aba, idx, codigo):
    """Dados gravados na linha idx do rascunho, ou None se a linha nao e (mais) desse codigo."""
    from gspread.utils import rowcol_to_a1
    valores = aba.get(f"A{idx}:{rowcol_to_a1(idx, len(COLUNAS_RASCUNHO) - 1 + BLOCOS_RASCUNHO)}")
    linha = valores[0] if valores else []
    if not linha or linha[0] != codigo:
        return None
    try:
        return json.loads("".join(linha[5:]))
    except Exception:
        return {}


def _juntar_anexo(dados, num, info):
    """Acrescenta `info` aos anexos do item `num` em `dados` (sem duplicar, respeitando remocoes)."""
    chave = _chave_anexo(info)
    if chave in (dados.get("anexos_removidos") or []):
        return False
    lista = dados.setdefault("anexos", {}).setdefault(str(num), [])
    if any(_chave_anexo(a) == chave for a in lista):
        return False
    lista.append(info)
    return True


def registrar_anexos_lote(codigo, pares):
    """Grava varios anexos [(num, info), ...] DIRETO no rascunho guardado, sem depender da sessao
    do navegador, numa unica escrita. Nunca lanca excecao. Devolve quantos foram acrescentados
    (0 se todos ja constavam/foram removidos) ou None se nao conseguiu gravar."""
    import time
    from gspread.utils import rowcol_to_a1
    try:
        with _trava(codigo):
            agora = agora_brasilia().strftime("%d/%m/%Y %H:%M:%S")
            if not _usa_google(_carregar_mapeamento()):
                caminho = _caminho_rascunho_local(codigo)
                if not os.path.exists(caminho):
                    return None
                with open(caminho, encoding="utf-8") as f:
                    d = json.load(f)
                novos = sum(1 for n, i in pares if _juntar_anexo(d["dados"], n, i))
                if novos:
                    d["atualizado_em"] = agora
                    with open(caminho, "w", encoding="utf-8") as f:
                        json.dump(d, f, ensure_ascii=False)
                return novos
            for tentativa in range(len(ESPERAS_429) + 1):
                try:
                    aba = _aba_rascunho()
                    idx = _linha_do_codigo(aba, codigo)
                    if not idx:
                        return None
                    cab = (aba.get(f"A{idx}:E{idx}") or [[]])[0] + [""] * 5
                    dados = _ler_dados_linha(aba, idx, codigo)
                    if not isinstance(dados, dict) or not dados:
                        return None
                    novos = sum(1 for n, i in pares if _juntar_anexo(dados, n, i))
                    if novos:
                        linha = _montar_linha(codigo, cab[1], cab[3], dados, agora)
                        aba.update(f"A{idx}:{rowcol_to_a1(idx, len(linha))}", [linha], value_input_option=RAW)
                        _CACHE_RASC["linhas"][codigo] = (idx, time.time())
                    return novos
                except Exception as exc:
                    _limpar_cache_rascunho()
                    if not _e_429(exc) or tentativa >= len(ESPERAS_429):
                        return None
                    time.sleep(ESPERAS_429[tentativa])
    except Exception:
        return None
    return None


def registrar_anexo_rascunho(codigo, num, info):
    """Um anexo (ver registrar_anexos_lote). True se a gravacao funcionou."""
    return registrar_anexos_lote(codigo, [(num, info)]) is not None


_PADRAO_ITEM = re.compile(r"^item(\d{1,3})_", re.I)


def recuperar_anexos_do_drive(codigo):
    """Religa ao rascunho as fotos/PDFs que estao na pasta da visita no Drive mas nao constam no
    rascunho (ex.: envio interrompido). O item vem do nome do arquivo (item{n}_...). Respeita os
    anexos removidos de proposito. Retorna {'na_pasta','novos','pares':[(num,info)],'erro'}."""
    from evidencias import EvidenciaErro, UploadAppsScript
    saida = {"na_pasta": 0, "novos": 0, "pares": [], "erro": ""}
    destino, _pasta = drive_para_evidencias()
    if not isinstance(destino, UploadAppsScript):
        saida["erro"] = "A recuperação precisa do serviço de upload do Drive (Apps Script) configurado."
        return saida
    try:
        arquivos = destino.listar(codigo)
    except EvidenciaErro as exc:
        saida["erro"] = str(exc)
        return saida
    except Exception:
        saida["erro"] = "Falha de conexão ao consultar o Drive. Tente novamente."
        return saida
    pares = []
    for a in sorted(arquivos, key=lambda x: (x.get("criado") or "", x.get("nome") or "")):
        m = _PADRAO_ITEM.match(a.get("nome") or "")
        mime = a.get("mime") or ""
        if not m or not (mime.startswith("image/") or mime == "application/pdf"):
            continue
        criado = a.get("criado") or ""
        try:
            from datetime import datetime, timezone, timedelta
            dt = datetime.fromisoformat(criado.replace("Z", "+00:00")).astimezone(timezone(timedelta(hours=-3)))
            enviado = dt.strftime("%d/%m/%Y %H:%M:%S")
        except Exception:
            enviado = ""
        pares.append((int(m.group(1)), {
            "nome": a["nome"], "tipo": "pdf" if mime == "application/pdf" else "imagem",
            "tamanho_kb": round((a.get("tamanho") or 0) / 1024, 1), "enviado_em": enviado,
            "link": a.get("link", ""), "id_arquivo": a.get("id", ""), "destino": "drive"}))
    saida["na_pasta"] = len(pares)
    if not pares:
        return saida
    novos = registrar_anexos_lote(codigo, pares)
    if novos is None:
        saida["erro"] = "Não foi possível gravar no rascunho agora. Tente novamente em instantes."
        return saida
    saida["novos"], saida["pares"] = novos, pares
    return saida


def mesclado_pendente(codigo):
    """Dados mesclados com o que outra sessao gravou (uma unica vez), ou None."""
    with _TRAVA_FILA:
        return _MESCLADO.pop(codigo, None)


# Controle de cota do Google Sheets (60 leituras/min por usuario de servico, compartilhadas
# entre todos os visitantes): aba e linha do rascunho ficam em cache, gravacoes rapidas sao
# agrupadas (so a ultima versao e gravada) e erros 429 sao repetidos com espera.
INTERVALO_MIN_GRAVACAO = 5.0          # segundos entre gravacoes do mesmo rascunho
ESPERAS_429 = (2, 5, 10, 20)          # segundos entre tentativas quando a cota estoura
VALIDADE_LINHA = 600                  # segundos ate reconferir a linha do rascunho
_CACHE_RASC = {"aba": None, "linhas": {}}
_PENDENTE = {}
_ESCRITOR = {}
_ULTIMA_GRAVACAO = {}
_TRAVA_FILA = threading.Lock()


def _limpar_cache_rascunho():
    _CACHE_RASC["aba"] = None
    _CACHE_RASC["linhas"] = {}


def _aba_rascunho():
    if _CACHE_RASC["aba"] is None:
        aba = _aba(_abrir_planilha(), ABA_RASCUNHO, COLUNAS_RASCUNHO)
        precisa = len(COLUNAS_RASCUNHO) - 1 + BLOCOS_RASCUNHO
        if getattr(aba, "col_count", precisa) < precisa:
            aba.add_cols(precisa - aba.col_count)
        _CACHE_RASC["aba"] = aba
    return _CACHE_RASC["aba"]


def _e_429(exc):
    txt = str(exc)
    return "429" in txt or "Quota exceeded" in txt or "RESOURCE_EXHAUSTED" in txt


def _montar_linha(codigo, hospital, responsavel, dados, agora):
    texto = json.dumps(dados, ensure_ascii=False)
    blocos = [texto[i:i + TAM_BLOCO_RASCUNHO] for i in range(0, len(texto), TAM_BLOCO_RASCUNHO)] or [""]
    if len(blocos) > BLOCOS_RASCUNHO:
        raise ValueError("Rascunho grande demais para o Google Sheets.")
    blocos += [""] * (BLOCOS_RASCUNHO - len(blocos))      # limpa blocos antigos
    return [codigo, hospital, chave_hospital(hospital), responsavel, agora] + blocos


def _gravar_rascunho_google(codigo, hospital, responsavel, dados, agora):
    """Le o que ja esta gravado, mescla com `dados` e grava. Devolve os dados efetivamente gravados."""
    import time
    from gspread.utils import rowcol_to_a1
    ultimo = None
    for tentativa in range(len(ESPERAS_429) + 1):
        try:
            aba = _aba_rascunho()
            cache = _CACHE_RASC["linhas"].get(codigo)
            idx = cache[0] if cache and time.time() - cache[1] < VALIDADE_LINHA else None
            guardado = None
            if idx is not None:
                guardado = _ler_dados_linha(aba, idx, codigo)
                if guardado is None:                   # a linha mudou de lugar (linhas apagadas/movidas)
                    idx = None
            if idx is None:
                idx = _linha_do_codigo(aba, codigo)
                if idx:
                    guardado = _ler_dados_linha(aba, idx, codigo)
            efetivo = mesclar_dados(guardado, dados) if guardado else dados
            linha = _montar_linha(codigo, hospital, responsavel, efetivo, agora)
            if idx:
                aba.update(f"A{idx}:{rowcol_to_a1(idx, len(linha))}", [linha], value_input_option=RAW)
                _CACHE_RASC["linhas"][codigo] = (idx, time.time())
            else:
                resp = aba.append_row(linha, value_input_option=RAW)
                try:
                    rng = resp["updates"]["updatedRange"]          # ex.: Rascunho!A7:F7
                    novo_idx = int("".join(c for c in rng.split("!")[1].split(":")[0] if c.isdigit()))
                    _CACHE_RASC["linhas"][codigo] = (novo_idx, time.time())
                except Exception:
                    _CACHE_RASC["linhas"].pop(codigo, None)
            return efetivo
        except Exception as exc:
            ultimo = exc
            _limpar_cache_rascunho()
            if not _e_429(exc) or tentativa >= len(ESPERAS_429):
                raise
            time.sleep(ESPERAS_429[tentativa])
    raise ultimo


def _gravar_agora(codigo, hospital, responsavel, dados, versao):
    with _trava(codigo):
        if versao < _ULTIMA_VERSAO.get(codigo, -1):
            return  # ja existe gravacao mais nova
        agora = agora_brasilia().strftime("%d/%m/%Y %H:%M:%S")
        try:
            mapping = _carregar_mapeamento()
            if _usa_google(mapping):
                efetivo = _gravar_rascunho_google(codigo, hospital, responsavel, dados, agora)
                if json.dumps(efetivo, sort_keys=True) != json.dumps(dados, sort_keys=True):
                    with _TRAVA_FILA:
                        _MESCLADO[codigo] = efetivo     # a sessao incorpora o que veio de fora
            else:
                os.makedirs(LOCAL_DIR, exist_ok=True)
                with open(_caminho_rascunho_local(codigo), "w", encoding="utf-8") as f:
                    json.dump({"codigo": codigo, "hospital": hospital, "responsavel": responsavel,
                               "atualizado_em": agora, "dados": dados}, f, ensure_ascii=False)
            _ULTIMA_VERSAO[codigo] = versao
            _ESTADO_AUTOSAVE[codigo] = {"ok": True, "quando": agora, "msg": ""}
        except Exception as exc:
            _ESTADO_AUTOSAVE[codigo] = {"ok": False, "quando": agora, "msg": str(exc)[:200]}


def salvar_rascunho(codigo, hospital, responsavel, dados, versao):
    """Autosave. Nunca lanca excecao; o resultado fica em estado_autosave().
    Chamadas em rajada sao agrupadas: so a versao mais nova e gravada, com intervalo minimo."""
    import time
    if not _usa_google(_carregar_mapeamento()):
        _gravar_agora(codigo, hospital, responsavel, dados, versao)   # modo local: sem limite de cota
        return
    with _TRAVA_FILA:
        atual = _PENDENTE.get(codigo)
        if atual is None or versao >= atual[3]:
            _PENDENTE[codigo] = (hospital, responsavel, dados, versao)
        if _ESCRITOR.get(codigo):
            return                      # outra thread ja esta gravando e pegara a mais nova
        _ESCRITOR[codigo] = True
    try:
        while True:
            with _TRAVA_FILA:
                item = _PENDENTE.pop(codigo, None)
                if item is None:
                    _ESCRITOR[codigo] = False
                    return
            espera = INTERVALO_MIN_GRAVACAO - (time.time() - _ULTIMA_GRAVACAO.get(codigo, 0))
            if espera > 0:
                time.sleep(espera)
                with _TRAVA_FILA:       # pode ter chegado versao mais nova durante a espera
                    novo = _PENDENTE.pop(codigo, None)
                    if novo is not None:
                        item = novo
            _ULTIMA_GRAVACAO[codigo] = time.time()
            _gravar_agora(codigo, item[0], item[1], item[2], item[3])
    except BaseException:
        with _TRAVA_FILA:
            _ESCRITOR[codigo] = False
        raise


def estado_autosave(codigo):
    """{'ok': bool, 'quando': str, 'msg': str} do ultimo autosave, ou None."""
    return _ESTADO_AUTOSAVE.get(codigo)


def _entrada(codigo, hospital, responsavel, atualizado, dados_json):
    try:
        dados = json.loads(dados_json)
        corrompido = False
    except Exception:
        dados, corrompido = None, True
    return {"codigo": codigo, "hospital": hospital, "responsavel": responsavel,
            "atualizado_em": atualizado, "dados": dados, "corrompido": corrompido}


def _todos_rascunhos():
    """Lista de entradas de TODOS os rascunhos. Levanta excecao se nao conseguir ler."""
    mapping = _carregar_mapeamento()
    if _usa_google(mapping):
        aba = _aba_se_existe(_abrir_planilha(), ABA_RASCUNHO)
        if aba is None:
            return []
        saida = []
        for linha in aba.get_all_values()[1:]:
            linha = linha + [""] * 6
            if linha[0]:
                saida.append(_entrada(linha[0], linha[1], linha[3], linha[4], "".join(linha[5:])))
        return saida
    saida = []
    if os.path.isdir(LOCAL_DIR):
        for nome in sorted(os.listdir(LOCAL_DIR)):
            if nome.startswith("Rascunho_") and nome.endswith(".json"):
                try:
                    with open(os.path.join(LOCAL_DIR, nome), encoding="utf-8") as f:
                        d = json.load(f)
                    saida.append({"codigo": d["codigo"], "hospital": d["hospital"], "responsavel": d["responsavel"],
                                  "atualizado_em": d["atualizado_em"], "dados": d["dados"], "corrompido": False})
                except Exception:
                    saida.append({"codigo": nome[9:-5], "hospital": "", "responsavel": "",
                                  "atualizado_em": "", "dados": None, "corrompido": True})
    return saida


def listar_rascunhos_abertos():
    """Todas as visitas em aberto (rascunhos), mais recentes primeiro, ou ERRO_VERIFICACAO_RASCUNHO."""
    try:
        def chave(r):
            try:
                return datetime.strptime(r["atualizado_em"], "%d/%m/%Y %H:%M:%S")
            except Exception:
                return datetime.min
        return sorted(_todos_rascunhos(), key=chave, reverse=True)
    except Exception:
        return ERRO_VERIFICACAO_RASCUNHO


def listar_rascunhos_hospital(hospital):
    """Lista (pode ser vazia = confirmado que nao ha) ou ERRO_VERIFICACAO_RASCUNHO."""
    try:
        chave = chave_hospital(hospital)
        return [r for r in _todos_rascunhos() if chave_hospital(r["hospital"]) == chave or r["corrompido"]]
    except Exception:
        return ERRO_VERIFICACAO_RASCUNHO


def carregar_rascunho_codigo(codigo):
    """Entrada, None (confirmado que nao existe) ou ERRO_VERIFICACAO_RASCUNHO."""
    try:
        codigo = (codigo or "").strip().upper()
        for r in _todos_rascunhos():
            if r["codigo"].upper() == codigo:
                return r
        return None
    except Exception:
        return ERRO_VERIFICACAO_RASCUNHO


def apagar_rascunho(codigo):
    """Chamado apos envio com sucesso. Nunca lanca excecao."""
    try:
        mapping = _carregar_mapeamento()
        if _usa_google(mapping):
            aba = _aba_se_existe(_abrir_planilha(), ABA_RASCUNHO)
            if aba is not None:
                idx = _linha_do_codigo(aba, codigo)
                if idx:
                    aba.delete_rows(idx)
        else:
            caminho = _caminho_rascunho_local(codigo)
            if os.path.exists(caminho):
                os.remove(caminho)
    except Exception:
        pass
    _limpar_cache_rascunho()
    with _TRAVA_FILA:
        _PENDENTE.pop(codigo, None)
    _ULTIMA_VERSAO.pop(codigo, None)
    _ESTADO_AUTOSAVE.pop(codigo, None)
    with _TRAVA_FILA:
        _MESCLADO.pop(codigo, None)


def modo_google():
    """True quando as credenciais do Google estao configuradas (senao o app grava so localmente)."""
    try:
        return bool(_usa_google(_carregar_mapeamento()))
    except Exception:
        return False
