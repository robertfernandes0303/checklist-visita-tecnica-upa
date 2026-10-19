import os, sys, shutil
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from streamlit.testing.v1 import AppTest
import schema as sc, sheets_backend as sb, evidencias as ev
APP = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "app.py")

def limpar():
    for d in (sb.LOCAL_DIR, os.path.join(os.path.dirname(APP), "evidencias_locais")):
        if os.path.isdir(d):
            for f in os.listdir(d):
                p = os.path.join(d, f)
                if f == ".gitkeep": continue
                shutil.rmtree(p) if os.path.isdir(p) else os.remove(p)

def test_fluxo_completo_incompleto():
    limpar()
    at = AppTest.from_file(APP, default_timeout=30).run()
    assert not at.exception
    at.text_input[0].set_value("UPA Teste Ação").run()
    at.text_input[1].set_value("Guajará-Mirim").run()
    at.selectbox[0].set_value("RO").run()
    at.text_input[2].set_value("Maria").run()
    [b for b in at.button if b.label == "Iniciar visita"][0].click().run()
    assert not at.exception and at.session_state.etapa == "menu"
    cod = at.session_state.codigo
    [b for b in at.button if b.key == "abrir_0"][0].click().run()
    assert at.session_state.etapa == "grupo"
    primeiro = sc.GRUPOS[0]["itens"][0]["num"]
    at.text_input[0].set_value("120").run()
    at.radio[0].set_value("Observado").run()
    at.text_area[0].set_value("obs teste").run()
    r = at.session_state.respostas[primeiro]
    assert r["resposta"] == "120" and r["fonte"] == "Observado" and r["obs"] == "obs teste"
    at.text_area[-1].set_value("impacto secao").run()
    [b for b in at.button if b.label == "Salvar e voltar ao menu"][0].click().run()
    import time; time.sleep(0.5)
    d = sb.carregar_rascunho_codigo(cod)["dados"]
    assert d["respostas"][str(primeiro)]["resposta"] == "120" and d["impactos"]["0"] == "impacto secao"
    [b for b in at.button if b.label == "Enviar mesmo incompleta"][0].click().run()
    [b for b in at.button if b.label == "Confirmar e enviar"][0].click().run()
    assert at.session_state.etapa == "confirmacao", [e.value for e in at.error]
    assert sb.carregar_rascunho_codigo(cod) is None
    import openpyxl
    wb = openpyxl.load_workbook(sb.ARQUIVO_LOCAL)
    assert wb[sb.ABA_VISITAS].max_row == 2 and wb[sb.ABA_ITENS].max_row == 1 + sc.TOTAL_ITENS
    h = [c.value for c in wb[sb.ABA_VISITAS][1]]
    row = [c.value for c in wb[sb.ABA_VISITAS][2]]
    assert row[h.index("Status da Visita")] == "Incompleta" and row[h.index("Cidade")] == "Guajará-Mirim"
    assert row[h.index("UF")] == "RO" and row[h.index("Total Verificados")] == "1"
    assert wb[sb.ABA_SECOES].max_row == 2

def test_retomar_por_hospital():
    limpar()
    sb.salvar_rascunho("ENF-X-1", "Hospital São José", "Ana", {"codigo_visita": "ENF-X-1", "hospital": "Hospital São José",
        "data_visita": "2026-10-08", "respostas": {"2": {"resposta": "x", "fonte": "Informado"}}}, 1)
    at = AppTest.from_file(APP, default_timeout=30).run()
    at.text_input[0].set_value("hospital sao jose").run()
    at.text_input[2].set_value("Ana").run()
    [b for b in at.button if b.label == "Iniciar visita"][0].click().run()
    assert at.session_state.etapa == "retomar_lista"
    [b for b in at.button if b.label == "Continuar"][0].click().run()
    assert at.session_state.etapa == "menu" and at.session_state.respostas[2]["resposta"] == "x"

def test_bloqueio_erro_verificacao(monkeypatch=None):
    limpar()
    orig = sb.listar_rascunhos_hospital
    sb.listar_rascunhos_hospital = lambda h: sb.ERRO_VERIFICACAO_RASCUNHO
    try:
        at = AppTest.from_file(APP, default_timeout=30).run()
        at.text_input[0].set_value("H").run(); at.text_input[2].set_value("R").run()
        [b for b in at.button if b.label == "Iniciar visita"][0].click().run()
        assert at.session_state.etapa == "bloqueio" and not at.session_state.codigo
    finally:
        sb.listar_rascunhos_hospital = orig

def test_evidencia_local_e_drive_mock():
    limpar()
    from PIL import Image; import io
    buf = io.BytesIO(); Image.new("RGB", (3000, 2000), "red").save(buf, "PNG")
    a = ev.guardar_evidencia(None, None, "ENF-T-1", 3, "foto.png", buf.getvalue())
    assert a["destino"] == "local" and a["tipo"] and os.path.exists(a["link"])
    try:
        ev.guardar_evidencia(None, None, "ENF-T-1", 3, "x.pdf", b"nao pdf")
        assert False
    except ev.EvidenciaErro: pass
    class R:
        def __init__(s, c, j): s.status_code=c; s._j=j; s.text=str(j); s.ok = c<400
        def json(s): return s._j
    chamadas=[]
    class Sess:
        def get(s,u,**k): chamadas.append(("GET",u)); return R(200,{"files":[]})
        def post(s,u,**k): chamadas.append(("POST",u,k.get("params"))); return R(200,{"id":"F1","webViewLink":"https://drive.google.com/x"})
    b = ev.guardar_evidencia(Sess(), "PASTA", "ENF-T-1", 3, "foto.png", buf.getvalue())
    assert b["destino"]=="drive" and b["link"].startswith("https://drive") and len(chamadas)>=2
    class Sess403(Sess):
        def post(s,u,**k): return R(403,{"error":{"errors":[{"reason":"storageQuotaExceeded"}],"message":"quota"}})
    try:
        ev.guardar_evidencia(Sess403(), "PASTA", "ENF-T-1", 3, "foto.png", buf.getvalue()); assert False
    except ev.EvidenciaErro as e: assert "Drive" in str(e) or "compart" in str(e).lower(), str(e)

def test_reenvio_idempotente():
    limpar()
    v={c:"" for c in sc.COLUNAS_VISITA}; v["Código da Visita"]="ENF-I-1"
    it=[{c:"" for c in sc.COLUNAS_ITEM}]; it[0]["Código da Visita"]="ENF-I-1"
    assert sb.salvar_auditoria(v,it,[])[0]
    assert sb.salvar_auditoria(v,it,[])[0]
    import openpyxl
    wb=openpyxl.load_workbook(sb.ARQUIVO_LOCAL)
    assert wb[sb.ABA_VISITAS].max_row==2 and wb[sb.ABA_ITENS].max_row==2

def test_versao_autosave():
    limpar()
    sb.salvar_rascunho("ENF-V-1","H","R",{"v":2},2)
    sb.salvar_rascunho("ENF-V-1","H","R",{"v":1},1)
    assert sb.carregar_rascunho_codigo("ENF-V-1")["dados"]["v"]==2

def test_lista_visitas_em_aberto_no_inicio():
    limpar()
    for i, h in enumerate(["Hospital A", "Hospital B"], 1):
        sb.salvar_rascunho(f"ENF-L-{i}", h, "Ana", {"codigo_visita": f"ENF-L-{i}", "hospital": h,
            "data_visita": "2026-10-08", "respostas": {"1": {"resposta": "9", "fonte": "Observado"}}}, 1)
    at = AppTest.from_file(APP, default_timeout=30).run()
    assert not at.exception
    botoes = [b for b in at.button if b.key and b.key.startswith("abrir_rasc_")]
    assert len(botoes) == 2
    botoes[1].click().run()
    assert at.session_state.etapa == "menu" and at.session_state.respostas[1]["resposta"] == "9"

def test_inicio_sem_abertas_e_falha_consulta():
    limpar()
    at = AppTest.from_file(APP, default_timeout=30).run()
    assert not [b for b in at.button if b.key and b.key.startswith("abrir_rasc_")]
    orig = sb.listar_rascunhos_abertos
    sb.listar_rascunhos_abertos = lambda: sb.ERRO_VERIFICACAO_RASCUNHO
    try:
        at = AppTest.from_file(APP, default_timeout=30).run()
        assert any("Não foi possível consultar" in w.value for w in at.warning)
    finally:
        sb.listar_rascunhos_abertos = orig


def _iniciar(at, hospital="Hospital <b>X</b> & Cia", resp="Ana"):
    at.text_input[0].set_value(hospital).run()
    at.text_input[2].set_value(resp).run()
    [b for b in at.button if b.label == "Iniciar visita"][0].click().run()
    assert at.session_state.etapa == "menu", [e.value for e in at.error]


def test_visita_completa_todos_itens_e_metricas():
    limpar()
    at = AppTest.from_file(APP, default_timeout=120).run()
    _iniciar(at)
    for gi, g in enumerate(sc.GRUPOS):
        [b for b in at.button if b.key == f"abrir_{gi}"][0].click().run()
        if gi == 0:
            [b for b in at.button if b.label == "Concluir seção"][0].click().run()
            assert at.session_state.etapa == "grupo" and any("Faltam" in w.value for w in at.warning)
        for r in at.radio:
            r.set_value("Não verificado" if gi % 2 else "Não se aplica")
        at.run()
        [b for b in at.button if b.label == "Concluir seção"][0].click().run()
        assert at.session_state.etapa == "menu", (gi, [w.value for w in at.warning])
    envio = [b for b in at.button if b.label == "Enviar visita completa"][0]
    assert not envio.disabled
    envio.click().run()
    [b for b in at.button if b.label == "Confirmar e enviar"][0].click().run()
    assert at.session_state.etapa == "confirmacao"
    import openpyxl
    wb = openpyxl.load_workbook(sb.ARQUIVO_LOCAL)
    h = [c.value for c in wb[sb.ABA_VISITAS][1]]
    v = [c.value for c in wb[sb.ABA_VISITAS][2]]
    assert v[h.index("Status da Visita")] == "Completa" and int(v[h.index("Total Não Preenchido")]) == 0
    assert int(v[h.index("Itens Respondidos")]) == sc.TOTAL_ITENS
    assert int(v[h.index("Total Não Verificados")]) + int(v[h.index("Total Não se Aplica")]) == sc.TOTAL_ITENS
    assert wb[sb.ABA_VISITAS].cell(2, 2).value == "Hospital <b>X</b> & Cia"
    [b for b in at.button if b.label == "Registrar nova visita"][0].click().run()
    assert at.session_state.etapa == "cabecalho" and not at.session_state.codigo and not at.session_state.respostas


def test_anexos_limite_e_remocao():
    limpar()
    from PIL import Image; import io
    at = AppTest.from_file(APP, default_timeout=30).run()
    _iniciar(at)
    [b for b in at.button if b.key == "abrir_0"][0].click().run()
    buf = io.BytesIO(); Image.new("RGB", (50, 50), "red").save(buf, "PNG")
    for _ in range(sc.MAX_ANEXOS_POR_ITEM):
        at.session_state.anexos.setdefault(1, []).append(
            ev.guardar_evidencia(None, None, at.session_state.codigo, 1, "f.png", buf.getvalue()))
    at.run()
    assert not at.exception
    rm = [b for b in at.button if b.key == "rm_1_0"]
    assert rm, "botao remover ausente"
    rm[0].click().run()
    assert len(at.session_state.anexos[1]) == sc.MAX_ANEXOS_POR_ITEM - 1


def test_nome_arquivo_com_acentos():
    assert ev.nome_seguro("Relatório Maçã.PDF") == "Relatorio_Maca.pdf"
    assert ev.nome_seguro("ç.pdf") == "c.pdf" and ev.nome_seguro("@@@.pdf") == "arquivo.pdf"
    assert ev.nome_seguro("../../x/y.png") == "y.png"


def test_formula_nao_executa_no_xlsx_local():
    limpar()
    v = {c: "" for c in sc.COLUNAS_VISITA}; v["Código da Visita"] = "ENF-F-1"; v["Unidade"] = "=1+1"
    it = [{c: "" for c in sc.COLUNAS_ITEM}]; it[0]["Código da Visita"] = "ENF-F-1"; it[0]["Observação"] = "=2+2"
    assert sb.salvar_auditoria(v, it, [])[0]
    import openpyxl
    wb = openpyxl.load_workbook(sb.ARQUIVO_LOCAL)
    assert wb[sb.ABA_VISITAS].cell(2, 2).data_type == "s" and wb[sb.ABA_ITENS].cell(2, sc.COLUNAS_ITEM.index("Observação") + 1).data_type == "s"


def test_upload_apps_script_mock_e_selecao():
    import json as _j
    from PIL import Image; import io
    buf = io.BytesIO(); Image.new("RGB", (80, 80), "red").save(buf, "PNG")
    os.environ["GOOGLE_UPLOAD_JSON"] = _j.dumps({"url": "https://script.google.com/macros/s/X/exec", "token": "tk"})
    orig = sb._carregar_mapeamento
    sb._carregar_mapeamento = lambda: {"pasta_evidencias": "PASTA123"}
    try:
        up, pasta = sb.drive_para_evidencias()
        assert isinstance(up, ev.UploadAppsScript) and pasta == "PASTA123"
    finally:
        sb._carregar_mapeamento = orig
        del os.environ["GOOGLE_UPLOAD_JSON"]
    import requests
    chamadas = []
    class R:
        def __init__(s, j): s._j = j
        def json(s): return s._j
    orig_post = requests.post
    requests.post = lambda url, data=None, **k: (chamadas.append((url, _j.loads(data))), R({"ok": True, "id": "F9", "link": "https://drive.google.com/file/d/F9"}))[1]
    try:
        r = ev.guardar_evidencia(ev.UploadAppsScript("https://x/exec", "tk"), "PASTA123", "ENF-U-1", 2, "foto.png", buf.getvalue())
        assert r["destino"] == "drive" and r["link"].endswith("F9")
        c = chamadas[0][1]
        assert c["token"] == "tk" and c["pasta"] == "ENF-U-1" and c["mime"] == "image/jpeg" and c["dados_b64"]
        requests.post = lambda url, data=None, **k: R({"ok": False, "erro": "token invalido"})
        try:
            ev.guardar_evidencia(ev.UploadAppsScript("https://x/exec", "tk"), "PASTA123", "ENF-U-1", 2, "foto.png", buf.getvalue()); assert False
        except ev.EvidenciaErro as e:
            assert "token invalido" in str(e)
    finally:
        requests.post = orig_post
